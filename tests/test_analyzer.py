"""Tests for the analyze command."""

import json

import pytest

from undatum.cmds.analyzer import OBJECTS_ANALYZE_LIMIT, Analyzer, analyze
from undatum.common.errors import FileNotFoundError


@pytest.fixture
def semicolon_csv(tmp_path):
    path = tmp_path / "orgs.csv"
    path.write_text(
        'id;name;city\n'
        '1;"Acme, Inc";"New York"\n'
        '2;"Beta LLC";London\n',
        encoding="utf8",
    )
    return str(path)


@pytest.fixture
def simple_csv(tmp_path):
    path = tmp_path / "people.csv"
    path.write_text(
        "id,name,city\n"
        "1,Alice,NYC\n"
        "2,Bob,NYC\n"
        "3,Carol,LA\n",
        encoding="utf8",
    )
    return str(path)


@pytest.fixture
def sample_xlsx(tmp_path):
    pytest.importorskip("openpyxl")
    from openpyxl import Workbook

    path = tmp_path / "sales.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["id", "amount"])
    ws.append(["1", "100"])
    ws.append(["2", "200"])
    ws.append(["3", "300"])
    wb.save(path)
    return str(path)


class TestAnalyzeCsv:
    def test_auto_detect_delimiter(self, semicolon_csv):
        report = analyze(semicolon_csv, engine="auto")
        assert report.success is True
        assert report.tables
        field_names = {f.name for f in report.tables[0].fields}
        assert field_names == {"id", "name", "city"}
        assert report.total_records == 2

    def test_explicit_delimiter(self, semicolon_csv):
        report = analyze(semicolon_csv, delimiter=";", engine="duckdb")
        field_names = {f.name for f in report.tables[0].fields}
        assert field_names == {"id", "name", "city"}

    def test_success_flag(self, simple_csv):
        report = analyze(simple_csv, engine="duckdb")
        assert report.success is True
        assert report.error is None

    def test_field_stats_populated(self, simple_csv):
        report = analyze(simple_csv, engine="duckdb", stats=True)
        name_field = next(f for f in report.tables[0].fields if f.name == "name")
        assert name_field.unique_count is not None
        assert name_field.total_count is not None
        assert name_field.uniqueness_pct is not None

    def test_no_stats(self, simple_csv):
        report = analyze(simple_csv, engine="duckdb", stats=False)
        for field in report.tables[0].fields:
            assert field.unique_count is None

    def test_no_scan_metadata_only(self, simple_csv):
        report = analyze(simple_csv, scan=False)
        assert report.success is True
        assert report.file_type == "csv"
        assert report.tables == []

    def test_analyzer_cli_wrapper_json(self, semicolon_csv, capsys):
        Analyzer().analyze(
            semicolon_csv,
            {
                "engine": "auto",
                "use_pandas": False,
                "outtype": "json",
                "output": None,
                "autodoc": False,
                "lang": "English",
                "delimiter": ";",
                "encoding": None,
                "objects_limit": OBJECTS_ANALYZE_LIMIT,
                "ignore_errors": True,
                "scan": True,
                "stats": True,
            },
        )
        captured = capsys.readouterr()
        payload = json.loads(captured.out)
        assert payload["success"] is True
        field_names = {f["name"] for f in payload["tables"][0]["fields"]}
        assert field_names == {"id", "name", "city"}

    def test_missing_file_raises(self, tmp_path):
        missing = str(tmp_path / "missing.csv")
        with pytest.raises(FileNotFoundError):
            Analyzer().analyze(missing, {"engine": "auto", "use_pandas": False, "outtype": "json"})

    def test_iterable_engine_csv(self, simple_csv):
        report = analyze(simple_csv, engine="iterable", stats=False)
        assert report.success is True
        assert report.total_records == 3
        assert len(report.tables) == 1
        field_names = {f.name for f in report.tables[0].fields}
        assert field_names == {"id", "name", "city"}

class TestAnalyzeXml:
    def test_nested_xml_records(self, tmp_path):
        xml_content = """<?xml version="1.0" encoding="UTF-8"?>
<Файл>
  <Документ ИдДок="a1" ДатаСост="10.12.2025">
    <СведМН><Регион>77</Регион></СведМН>
  </Документ>
  <Документ ИдДок="a2" ДатаСост="11.12.2025">
    <СведМН><Регион>50</Регион></СведМН>
  </Документ>
</Файл>"""
        path = tmp_path / "sample.xml"
        path.write_text(xml_content, encoding="utf8")
        report = analyze(str(path), engine="iterable", stats=False)
        assert report.success is True
        assert report.total_records == 2
        assert report.tables[0].num_cols > 0
        field_names = {f.name for f in report.tables[0].fields}
        assert "@ИдДок" in field_names or "ИдДок" in field_names or any("ИдДок" in n for n in field_names)


class TestAnalyzeXlsx:
    def test_reads_all_rows(self, sample_xlsx):
        report = analyze(sample_xlsx, engine="iterable", stats=False)
        assert report.success is True
        assert len(report.tables) == 1
        assert report.tables[0].num_records == 4
        assert report.tables[0].num_cols >= 2
