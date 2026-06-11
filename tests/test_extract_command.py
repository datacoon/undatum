"""Tests for extract command."""

from pathlib import Path

import pandas as pd
import pytest

from undatum.cmds.extractor import Extractor


def _find_output_file(output_dir: Path, suffix: str) -> Path:
    matches = list(output_dir.glob(f"*{suffix}"))
    assert matches, f"No output files with suffix {suffix}"
    return matches[0]


def test_extract_docx_tables_to_csv(tmp_path):
    pytest.importorskip("docx", reason="python-docx not installed")
    from docx import Document  # type: ignore

    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "name"
    table.cell(0, 1).text = "age"
    table.cell(1, 0).text = "Alice"
    table.cell(1, 1).text = "30"
    doc_path = tmp_path / "sample.docx"
    doc.save(doc_path)

    output_file = tmp_path / "out.csv"
    Extractor().extract(
        [str(doc_path)],
        {"output_format": "csv", "output": str(output_file), "method": "tables"},
    )

    content = output_file.read_text().strip().splitlines()
    assert content[0] == "name,age"
    assert "Alice,30" in content[1]


def test_extract_xlsx_to_ndjson_dir(tmp_path):
    df = pd.DataFrame({"city": ["Paris", "Berlin"], "population": [2, 4]})
    xlsx_path = tmp_path / "sample.xlsx"
    df.to_excel(xlsx_path, index=False)

    output_dir = tmp_path / "out"
    Extractor().extract(
        [str(xlsx_path)],
        {"output_format": "ndjson", "output_dir": str(output_dir)},
    )

    output_file = _find_output_file(output_dir, ".ndjson")
    content = output_file.read_text().strip().splitlines()
    assert len(content) == 2
