"""Tests for iterabledata Wave 1 CLI/SDK plumbing."""

import inspect
import json
import sqlite3
from pathlib import Path

import pytest
from typer.testing import CliRunner

from undatum.cli.data_commands import convert as convert_cmd
from undatum.cmds.converter import Converter, _skip_duckdb_convert
from undatum.cmds.head import Head
from undatum.cmds.schemer import build_schema
from undatum.cmds.statistics import StatProcessor
from undatum.common.command_utils import (
    apply_table_selection,
    iter_projected_rows,
    parse_column_list,
    parse_row_range,
    validate_codec_profile,
)
from undatum.common.errors import ValidationError
from undatum.core import app
from undatum.sdk.dataset import Dataset

runner = CliRunner()

NESTED_JSONL = (
    '{"name": "TJK", "capital_city": {"lat": 38.56, "lon": 68.77}}\n'
    '{"name": "KGZ", "capital_city": {"lat": 42.87, "lon": 74.59}}\n'
)


def _write_nested_jsonl(path: Path) -> Path:
    path.write_text(NESTED_JSONL, encoding="utf8")
    return path


def _write_two_sheet_xlsx(path: Path) -> Path:
    pytest.importorskip("openpyxl")
    from openpyxl import Workbook

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Sheet1"
    ws1.append(["id", "name"])
    ws1.append([1, "Alice"])
    ws2 = wb.create_sheet("Sheet2")
    ws2.append(["city", "pop"])
    ws2.append(["Dushanbe", 863400])
    wb.save(path)
    return path


def _write_two_table_sqlite(path: Path) -> Path:
    conn = sqlite3.connect(path)
    try:
        conn.execute("CREATE TABLE people (id INTEGER, name TEXT)")
        conn.execute("INSERT INTO people VALUES (1, 'Alice')")
        conn.execute("CREATE TABLE cities (city TEXT, pop INTEGER)")
        conn.execute("INSERT INTO cities VALUES ('Dushanbe', 863400)")
        conn.commit()
    finally:
        conn.close()
    return path


class TestHelpers:
    def test_parse_column_list(self):
        assert parse_column_list("id, name, city") == ["id", "name", "city"]
        assert parse_column_list("") is None
        assert parse_column_list(None) is None

    def test_parse_row_range(self):
        assert parse_row_range("0:1000") == (0, 1000)
        assert parse_row_range(None) is None
        with pytest.raises(ValidationError, match="row_range"):
            parse_row_range("1000")
        with pytest.raises(ValidationError, match="row_range"):
            parse_row_range("10:1")

    def test_validate_codec_profile(self):
        assert validate_codec_profile("MAX") == "max"
        assert validate_codec_profile(None) is None
        with pytest.raises(ValidationError, match="profile"):
            validate_codec_profile("turbo")

    def test_iter_projected_rows_flattens_dicts(self):
        rows = [{"capital_city": {"lat": 38.56}}]
        projected = list(iter_projected_rows(rows, True))
        assert projected[0]["capital_city.lat"] == 38.56
        unchanged = list(iter_projected_rows(rows, False))
        assert unchanged[0] == rows[0]

    def test_iter_projected_rows_respects_max_depth(self):
        rows = [{"a": {"b": {"c": 1}}}]
        deep = list(iter_projected_rows(rows, True))
        assert deep[0]["a.b.c"] == 1
        shallow = list(iter_projected_rows(rows, True, max_depth=1))
        assert "a.b.c" not in shallow[0]
        assert "a.b" in shallow[0]


class TestApplyTableSelection:
    def test_excel_sheet_name_becomes_page(self, tmp_path):
        xlsx = _write_two_sheet_xlsx(tmp_path / "wb.xlsx")
        args = apply_table_selection(str(xlsx), {"table": "Sheet2"})
        assert args["page"] == 1
        assert "table" not in args

    def test_excel_unknown_sheet_suggests(self, tmp_path):
        xlsx = _write_two_sheet_xlsx(tmp_path / "wb.xlsx")
        with pytest.raises(ValidationError, match="table") as exc:
            apply_table_selection(str(xlsx), {"table": "Missing"})
        assert "Sheet1" in str(exc.value)
        assert "Sheet2" in str(exc.value)

    def test_sqlite_keeps_table_name(self, tmp_path):
        db = _write_two_table_sqlite(tmp_path / "data.sqlite")
        args = apply_table_selection(str(db), {"table": "cities"})
        assert args["table"] == "cities"
        assert "page" not in args

    def test_uri_passes_table_through(self):
        args = apply_table_selection("s3://bucket/data.parquet", {"table": "events"})
        assert args["table"] == "events"

    def test_no_table_is_noop(self):
        assert apply_table_selection("x.csv", {"encoding": "utf8"}) == {"encoding": "utf8"}

    def test_force_iterable_if_table(self):
        from undatum.common.command_utils import force_iterable_if_table

        assert force_iterable_if_table({"table": "Sheet2"}, "duckdb") == "iterable"
        assert force_iterable_if_table({"sheet": "Sheet2"}, "duckdb") == "iterable"
        assert force_iterable_if_table({"table2": "Sheet2"}, "duckdb") == "iterable"
        assert force_iterable_if_table({"on_error": "skip"}, "duckdb") == "iterable"
        assert force_iterable_if_table({"on_error": "warn"}, "duckdb") == "iterable"
        assert force_iterable_if_table({"on_error": "raise"}, "duckdb") == "duckdb"
        assert force_iterable_if_table({"error_log": "errors.jsonl"}, "duckdb") == "iterable"
        assert force_iterable_if_table({"quotechar": "'"}, "duckdb") == "iterable"
        assert force_iterable_if_table({"flatten_nested": True}, "duckdb") == "iterable"
        assert force_iterable_if_table({}, "duckdb") == "duckdb"

    def test_get_side_iterable_options_does_not_inherit_table(self):
        from undatum.common.command_utils import get_side_iterable_options

        options = {"table": "Sheet2", "encoding": "utf8"}
        left = get_side_iterable_options(options, 1)
        right = get_side_iterable_options(options, 2)
        assert left["table"] == "Sheet2"
        assert "table" not in right
        both = get_side_iterable_options({"table": "A", "table2": "B"}, 2)
        assert both["table"] == "B"


class TestSkipDuckdbConvert:
    def test_skips_when_table_or_profile_or_native(self):
        assert _skip_duckdb_convert({"table": "Sheet2"}) is True
        assert _skip_duckdb_convert({"profile": "max"}) is True
        assert _skip_duckdb_convert({"level": 9}) is True
        assert _skip_duckdb_convert({"level": 0}) is True
        assert _skip_duckdb_convert({"native_batch": True}) is True
        assert _skip_duckdb_convert({"columns": "id"}) is True
        assert _skip_duckdb_convert({"write_mode": "overwrite"}) is True
        assert _skip_duckdb_convert({"on_error": "skip"}) is True
        assert _skip_duckdb_convert({"on_error": "warn"}) is True
        assert _skip_duckdb_convert({"on_error": "raise"}) is False
        assert _skip_duckdb_convert({"error_log": "errors.jsonl"}) is True
        assert _skip_duckdb_convert({"quotechar": "'"}) is True
        assert _skip_duckdb_convert({"row_group_size": 1024}) is True
        assert _skip_duckdb_convert({"progress": False}) is False


class TestOnErrorPolicy:
    def test_validate_on_error(self):
        from undatum.common.command_utils import get_iterable_options, validate_on_error

        assert validate_on_error(None) is None
        assert validate_on_error("SKIP") == "skip"
        assert get_iterable_options({"on_error": "warn"}) == {"on_error": "warn"}
        with pytest.raises(ValidationError, match="on-error"):
            validate_on_error("ignore")

    def test_convert_skips_bad_jsonl_rows(self, tmp_path):
        src = tmp_path / "mixed.jsonl"
        src.write_text('{"id": 1}\nnot-json\n{"id": 2}\n', encoding="utf8")
        dst = tmp_path / "out.jsonl"
        Converter().convert(
            str(src),
            str(dst),
            {"on_error": "skip", "progress": False, "engine": "iterable"},
        )
        rows = [json.loads(line) for line in dst.read_text(encoding="utf8").splitlines() if line]
        assert [row["id"] for row in rows] == [1, 2]

    def test_convert_writes_error_log(self, tmp_path):
        src = tmp_path / "mixed.jsonl"
        src.write_text('{"id": 1}\nnot-json\n{"id": 2}\n', encoding="utf8")
        dst = tmp_path / "out.jsonl"
        log = tmp_path / "errors.jsonl"
        Converter().convert(
            str(src),
            str(dst),
            {
                "on_error": "skip",
                "error_log": str(log),
                "progress": False,
                "engine": "iterable",
            },
        )
        entries = [json.loads(line) for line in log.read_text(encoding="utf8").splitlines() if line]
        assert entries
        assert any("not-json" in str(entry.get("original_line")) for entry in entries)


class TestFlattenNested:
    def test_schema_unfolds_dotted_fields(self, tmp_path):
        src = _write_nested_jsonl(tmp_path / "nested.jsonl")
        table = build_schema(str(src), options={"flatten_nested": True, "engine": "iterable"})
        names = [field.name for field in table.fields]
        assert "capital_city.lat" in names
        assert "capital_city.lon" in names
        assert "capital_city" not in names

    def test_schema_keep_nested_parents(self, tmp_path):
        src = _write_nested_jsonl(tmp_path / "nested.jsonl")
        table = build_schema(
            str(src),
            options={
                "flatten_nested": True,
                "keep_nested_parents": True,
                "engine": "iterable",
            },
        )
        names = [field.name for field in table.fields]
        assert "capital_city.lat" in names
        assert "capital_city" in names

    def test_schema_max_nested_depth(self, tmp_path):
        src = tmp_path / "deep.jsonl"
        src.write_text('{"a": {"b": {"c": 1}}}\n', encoding="utf8")
        deep = build_schema(str(src), options={"flatten_nested": True, "engine": "iterable"})
        shallow = build_schema(
            str(src),
            options={"flatten_nested": True, "max_nested_depth": 1, "engine": "iterable"},
        )
        deep_names = [field.name for field in deep.fields]
        shallow_names = [field.name for field in shallow.fields]
        assert "a.b.c" in deep_names
        assert "a.b.c" not in shallow_names

    def test_stats_unfolds_dotted_fields(self, tmp_path):
        src = _write_nested_jsonl(tmp_path / "nested.jsonl")
        out = tmp_path / "stats.json"
        profile = StatProcessor().stats(
            str(src),
            {
                "flatten_nested": True,
                "engine": "iterable",
                "format_out": "json",
                "output": str(out),
                "quiet": True,
                "progress": False,
            },
        )
        keys = {fd.get("key") for fd in (profile.get("debug") or {}).get("fielddata", {}).values()}
        assert "capital_city.lat" in keys
        assert "capital_city.lon" in keys

    def test_head_unfolds_dotted_fields(self, tmp_path, capsys):
        src = _write_nested_jsonl(tmp_path / "nested.jsonl")
        Head().head(str(src), {"flatten_nested": True, "n": 1})
        captured = capsys.readouterr().out
        row = json.loads(captured.splitlines()[0])
        assert row["capital_city.lat"] == 38.56
        assert row["name"] == "TJK"

    def test_select_flatten_uses_literal_dotted_keys(self, tmp_path):
        from undatum.cmds.selector import Selector

        src = _write_nested_jsonl(tmp_path / "nested.jsonl")
        dst = tmp_path / "out.jsonl"
        Selector().select(
            str(src),
            {
                "fields": "name,capital_city.lat",
                "flatten_nested": True,
                "output": str(dst),
                "format_out": "jsonl",
            },
        )
        row = json.loads(dst.read_text(encoding="utf8").splitlines()[0])
        assert row["name"] == "TJK"
        assert row["capital_city.lat"] == 38.56
        assert "capital_city" not in row

    def test_analyze_unfolds_dotted_fields(self, tmp_path):
        from undatum.cmds.analyzer import analyze

        src = _write_nested_jsonl(tmp_path / "nested.jsonl")
        report = analyze(
            str(src),
            engine="iterable",
            stats=False,
            options={"flatten_nested": True},
        )
        names = {field.name for field in report.tables[0].fields}
        assert "capital_city.lat" in names
        assert "capital_city.lon" in names

    def test_uniq_flatten_literal_dotted_key(self, tmp_path, capsys):
        from undatum.cmds.selector import Selector

        src = _write_nested_jsonl(tmp_path / "nested.jsonl")
        Selector().uniq(
            str(src),
            {"fields": "capital_city.lat", "flatten_nested": True, "engine": "iterable"},
        )
        captured = capsys.readouterr().out
        assert "38.56" in captured
        assert "42.87" in captured

    def test_headers_flatten_lists_dotted_keys(self, tmp_path, capsys):
        from undatum.cmds.selector import Selector

        src = _write_nested_jsonl(tmp_path / "nested.jsonl")
        Selector().headers(str(src), {"flatten_nested": True})
        captured = capsys.readouterr().out
        assert "capital_city.lat" in captured
        assert "capital_city.lon" in captured

    def test_sort_flatten_by_dotted_key(self, tmp_path):
        from undatum.cmds.sorter import Sorter

        src = _write_nested_jsonl(tmp_path / "nested.jsonl")
        dst = tmp_path / "sorted.jsonl"
        Sorter().sort(
            str(src),
            {
                "by": "capital_city.lat",
                "flatten_nested": True,
                "numeric": "capital_city.lat",
                "engine": "iterable",
                "output": str(dst),
            },
        )
        rows = [json.loads(line) for line in dst.read_text(encoding="utf8").splitlines() if line]
        assert [row["name"] for row in rows] == ["TJK", "KGZ"]
        assert rows[0]["capital_city.lat"] == 38.56

    def test_sniff_unfolds_dotted_fields(self, tmp_path, capsys):
        from undatum.cmds.sniffer import Sniffer

        src = _write_nested_jsonl(tmp_path / "nested.jsonl")
        Sniffer().sniff(str(src), {"flatten_nested": True, "format": "json"})
        captured = json.loads(capsys.readouterr().out)
        assert "capital_city.lat" in captured["fields"]
        assert "capital_city.lon" in captured["fields"]

    def test_split_flatten_by_dotted_key(self, tmp_path):
        from undatum.cmds.selector import Selector

        src = _write_nested_jsonl(tmp_path / "nested.jsonl")
        outdir = tmp_path / "parts"
        outdir.mkdir()
        Selector().split(
            str(src),
            {
                "fields": "capital_city.lat",
                "flatten_nested": True,
                "dirname": str(outdir),
                "zipfile": False,
            },
        )
        files = sorted(p.name for p in outdir.iterdir())
        assert "38.56.jsonl" in files
        assert "42.87.jsonl" in files
        row = json.loads((outdir / "38.56.jsonl").read_text(encoding="utf8").splitlines()[0])
        assert row["name"] == "TJK"
        assert row["capital_city.lat"] == 38.56

    def test_join_flatten_on_dotted_key(self, tmp_path):
        from undatum.cmds.joiner import Joiner

        left = _write_nested_jsonl(tmp_path / "left.jsonl")
        right = tmp_path / "right.jsonl"
        right.write_text(
            '{"capital_city": {"lat": 38.56}, "country": "Tajikistan"}\n'
            '{"capital_city": {"lat": 42.87}, "country": "Kyrgyzstan"}\n',
            encoding="utf8",
        )
        dst = tmp_path / "out.jsonl"
        Joiner().join(
            str(left),
            str(right),
            {
                "on": "capital_city.lat",
                "type": "inner",
                "flatten_nested": True,
                "engine": "iterable",
                "output": str(dst),
            },
        )
        rows = [json.loads(line) for line in dst.read_text(encoding="utf8").splitlines() if line]
        by_name = {row["name"]: row for row in rows}
        assert by_name["TJK"]["country"] == "Tajikistan"
        assert by_name["KGZ"]["country"] == "Kyrgyzstan"
        assert by_name["TJK"]["capital_city.lat"] == 38.56

    def test_exclude_flatten_on_dotted_key(self, tmp_path):
        from undatum.cmds.excluder import Excluder

        src = _write_nested_jsonl(tmp_path / "nested.jsonl")
        excl = tmp_path / "skip.jsonl"
        excl.write_text('{"capital_city": {"lat": 38.56}}\n', encoding="utf8")
        dst = tmp_path / "out.jsonl"
        Excluder().exclude(
            str(src),
            str(excl),
            {
                "on": "capital_city.lat",
                "flatten_nested": True,
                "output": str(dst),
            },
        )
        rows = [json.loads(line) for line in dst.read_text(encoding="utf8").splitlines() if line]
        assert [row["name"] for row in rows] == ["KGZ"]
        assert rows[0]["capital_city.lat"] == 42.87

    def test_diff_flatten_on_dotted_key(self, tmp_path):
        from undatum.cmds.differ import Differ

        left = _write_nested_jsonl(tmp_path / "left.jsonl")
        right = tmp_path / "right.jsonl"
        right.write_text(
            '{"name": "TJK", "capital_city": {"lat": 38.56, "lon": 68.77}}\n'
            '{"name": "KGZ", "capital_city": {"lat": 42.87, "lon": 99.0}}\n',
            encoding="utf8",
        )
        dst = tmp_path / "diff.json"
        Differ().diff(
            str(left),
            str(right),
            {
                "key": "name",
                "flatten_nested": True,
                "output_format": "json",
                "output": str(dst),
            },
        )
        captured = json.loads(dst.read_text(encoding="utf8"))
        assert captured["summary"]["changed_count"] == 1
        blob = json.dumps(captured)
        assert "capital_city.lon" in blob
        assert "99.0" in blob

    def test_validate_flatten_dotted_field(self, tmp_path, capsys):
        from undatum.cmds.validator import Validator

        src = _write_nested_jsonl(tmp_path / "nested.jsonl")
        rules = tmp_path / "rules.yml"
        rules.write_text(
            "rules:\n" "  - field: capital_city.lat\n" "    type: field\n" "    min: 40\n",
            encoding="utf8",
        )
        with pytest.raises(SystemExit) as exc:
            Validator().validate(
                str(src),
                {
                    "rules": str(rules),
                    "flatten_nested": True,
                    "output_format": "json",
                },
            )
        assert exc.value.code == 1
        captured = json.loads(capsys.readouterr().out)
        stats = captured["statistics"]
        assert stats["total_records"] == 2
        assert stats["errors"] == 1
        assert stats["passed"] == 1


class TestConvertTableSelection:
    def test_convert_xlsx_named_sheet(self, tmp_path):
        xlsx = _write_two_sheet_xlsx(tmp_path / "wb.xlsx")
        dst = tmp_path / "out.jsonl"
        Converter().convert(
            str(xlsx),
            str(dst),
            {"table": "Sheet2", "progress": False, "engine": "python"},
        )
        rows = [json.loads(line) for line in dst.read_text(encoding="utf8").splitlines() if line]
        assert rows
        assert "city" in rows[0]
        assert rows[0]["city"] == "Dushanbe"

    def test_convert_sqlite_named_table(self, tmp_path):
        db = _write_two_table_sqlite(tmp_path / "data.sqlite")
        dst = tmp_path / "out.jsonl"
        Converter().convert(
            str(db),
            str(dst),
            {"table": "cities", "progress": False, "engine": "python"},
        )
        rows = [json.loads(line) for line in dst.read_text(encoding="utf8").splitlines() if line]
        assert rows
        assert rows[0]["city"] == "Dushanbe"

    def test_head_xlsx_named_sheet(self, tmp_path, capsys):
        xlsx = _write_two_sheet_xlsx(tmp_path / "wb.xlsx")
        Head().head(str(xlsx), {"table": "Sheet2", "n": 5})
        captured = capsys.readouterr().out
        assert "Dushanbe" in captured
        assert "Alice" not in captured

    def test_tail_xlsx_named_sheet(self, tmp_path, capsys):
        from undatum.cmds.tail import Tail

        xlsx = _write_two_sheet_xlsx(tmp_path / "wb.xlsx")
        Tail().tail(str(xlsx), {"table": "Sheet2", "n": 5})
        captured = capsys.readouterr().out
        assert "Dushanbe" in captured
        assert "Alice" not in captured

    def test_count_xlsx_named_sheet(self, tmp_path, capsys):
        from undatum.cmds.counter import Counter

        xlsx = _write_two_sheet_xlsx(tmp_path / "wb.xlsx")
        Counter().count(str(xlsx), {"table": "Sheet2"})
        captured = capsys.readouterr().out.strip()
        assert captured == "1"

    def test_analyze_xlsx_named_sheet(self, tmp_path):
        from undatum.cmds.analyzer import analyze

        xlsx = _write_two_sheet_xlsx(tmp_path / "wb.xlsx")
        report = analyze(str(xlsx), table_name="Sheet2", stats=False)
        assert report.success is True
        assert len(report.tables) == 1
        assert report.tables[0].id == "Sheet2"
        field_names = {f.name for f in report.tables[0].fields}
        assert "city" in field_names

    def test_headers_xlsx_named_sheet(self, tmp_path, capsys):
        from undatum.cmds.selector import Selector

        xlsx = _write_two_sheet_xlsx(tmp_path / "wb.xlsx")
        Selector().headers(str(xlsx), {"table": "Sheet2"})
        captured = capsys.readouterr().out
        assert "city" in captured
        assert "id" not in captured

    def test_uniq_xlsx_named_sheet(self, tmp_path, capsys):
        from undatum.cmds.selector import Selector

        xlsx = _write_two_sheet_xlsx(tmp_path / "wb.xlsx")
        Selector().uniq(str(xlsx), {"table": "Sheet2", "fields": "city"})
        captured = capsys.readouterr().out
        assert "Dushanbe" in captured
        assert "Alice" not in captured

    def test_frequency_xlsx_named_sheet(self, tmp_path, capsys):
        from undatum.cmds.selector import Selector

        xlsx = _write_two_sheet_xlsx(tmp_path / "wb.xlsx")
        Selector().frequency(str(xlsx), {"table": "Sheet2", "fields": "city"})
        captured = capsys.readouterr().out
        assert "Dushanbe" in captured
        assert "Alice" not in captured

    def test_table_xlsx_named_sheet(self, tmp_path, capsys):
        from undatum.cmds.table import TableFormatter

        xlsx = _write_two_sheet_xlsx(tmp_path / "wb.xlsx")
        TableFormatter().table(str(xlsx), {"table": "Sheet2", "limit": 5})
        captured = capsys.readouterr().out
        assert "Dushanbe" in captured
        assert "Alice" not in captured

    def test_sort_xlsx_named_sheet(self, tmp_path):
        from undatum.cmds.sorter import Sorter

        xlsx = _write_two_sheet_xlsx(tmp_path / "wb.xlsx")
        dst = tmp_path / "out.jsonl"
        Sorter().sort(
            str(xlsx),
            {"table": "Sheet2", "by": "city", "output": str(dst), "engine": "iterable"},
        )
        text = dst.read_text(encoding="utf8")
        assert "Dushanbe" in text
        assert "Alice" not in text

    def test_mask_xlsx_named_sheet(self, tmp_path):
        from undatum.cmds.masker import Masker

        xlsx = _write_two_sheet_xlsx(tmp_path / "wb.xlsx")
        dst = tmp_path / "out.jsonl"
        Masker().mask(
            str(xlsx),
            str(dst),
            {"table": "Sheet2", "fields": "city", "method": "redact"},
        )
        text = dst.read_text(encoding="utf8")
        assert "Alice" not in text
        assert "***" in text or "Dushanbe" not in text

    def test_dataset_sort_preserves_table(self, tmp_path):
        xlsx = _write_two_sheet_xlsx(tmp_path / "wb.xlsx")
        rows = list(Dataset.read(str(xlsx), table="Sheet2").sort("city"))
        assert rows
        assert "city" in rows[0]
        assert rows[0]["city"] == "Dushanbe"

    def test_rename_xlsx_named_sheet(self, tmp_path):
        from undatum.cmds.renamer import Renamer

        xlsx = _write_two_sheet_xlsx(tmp_path / "wb.xlsx")
        dst = tmp_path / "out.jsonl"
        Renamer().rename(
            str(xlsx),
            {"table": "Sheet2", "map": "city:city_name", "output": str(dst)},
        )
        text = dst.read_text(encoding="utf8")
        assert "city_name" in text
        assert "Dushanbe" in text
        assert "Alice" not in text

    def test_enum_xlsx_named_sheet(self, tmp_path):
        from undatum.cmds.enumerator import Enumerator

        xlsx = _write_two_sheet_xlsx(tmp_path / "wb.xlsx")
        dst = tmp_path / "out.jsonl"
        Enumerator().enum(
            str(xlsx),
            {"table": "Sheet2", "field": "row_id", "output": str(dst)},
        )
        text = dst.read_text(encoding="utf8")
        assert "Dushanbe" in text
        assert "Alice" not in text
        assert "row_id" in text

    def test_join_xlsx_named_sheet(self, tmp_path):
        from undatum.cmds.joiner import Joiner

        xlsx = _write_two_sheet_xlsx(tmp_path / "wb.xlsx")
        right = tmp_path / "right.csv"
        right.write_text("city,country\nDushanbe,TJK\n", encoding="utf8")
        dst = tmp_path / "out.jsonl"
        Joiner().join(
            str(xlsx),
            str(right),
            {
                "table": "Sheet2",
                "on": "city",
                "type": "inner",
                "output": str(dst),
                "engine": "iterable",
            },
        )
        text = dst.read_text(encoding="utf8")
        assert "Dushanbe" in text
        assert "TJK" in text
        assert "Alice" not in text

    def test_exclude_xlsx_named_sheet(self, tmp_path):
        from undatum.cmds.excluder import Excluder

        xlsx = _write_two_sheet_xlsx(tmp_path / "wb.xlsx")
        excl = tmp_path / "excl.csv"
        excl.write_text("city\nDushanbe\n", encoding="utf8")
        dst = tmp_path / "out.jsonl"
        Excluder().exclude(
            str(xlsx),
            str(excl),
            {"table": "Sheet2", "on": "city", "output": str(dst)},
        )
        text = dst.read_text(encoding="utf8")
        assert "Alice" not in text
        assert "Dushanbe" not in text

    def test_package_xlsx_named_sheet(self, tmp_path):
        from undatum.cmds.packager import Packager

        xlsx = _write_two_sheet_xlsx(tmp_path / "wb.xlsx")
        dst = tmp_path / "datapackage.json"
        result = Packager().create(
            [str(xlsx)],
            {"output": str(dst), "table": "Sheet2", "quiet": True, "engine": "iterable"},
        )
        fields = [
            field["name"]
            for resource in result["package"]["resources"]
            for field in resource.get("schema", {}).get("fields", [])
        ]
        assert "city" in fields
        assert "name" not in fields

    def test_dataset_rename_preserves_table(self, tmp_path):
        xlsx = _write_two_sheet_xlsx(tmp_path / "wb.xlsx")
        rows = list(Dataset.read(str(xlsx), table="Sheet2").rename({"city": "city_name"}))
        assert rows
        assert "city_name" in rows[0]
        assert rows[0]["city_name"] == "Dushanbe"
        assert "Alice" not in str(rows)

    def test_split_xlsx_named_sheet(self, tmp_path):
        from undatum.cmds.selector import Selector

        xlsx = _write_two_sheet_xlsx(tmp_path / "wb.xlsx")
        Selector().split(
            str(xlsx),
            {
                "table": "Sheet2",
                "fields": "city",
                "dirname": str(tmp_path),
                "chunksize": 10000,
                "filter": None,
            },
        )
        out = tmp_path / "Dushanbe.jsonl"
        assert out.exists()
        text = out.read_text(encoding="utf8")
        assert "Dushanbe" in text
        assert "Alice" not in text
        assert not (tmp_path / "Alice.jsonl").exists()

    def test_tool_frequency_xlsx_named_sheet(self, tmp_path):
        from undatum.tools import frequency

        xlsx = _write_two_sheet_xlsx(tmp_path / "wb.xlsx")
        result = frequency(str(xlsx), "city", table="Sheet2")
        assert result["ok"] is True
        assert result["data"]["total_rows"] == 1
        assert result["data"]["top_values"][0]["value"] == "Dushanbe"
        from undatum.tools import schemas as tool_schemas

        freq_schema = next(
            d for d in tool_schemas.UNDATUM_TOOL_DEFINITIONS if d["name"] == "frequency"
        )
        assert "table" in freq_schema["parameters"]["properties"]

    def test_tool_frequency_flatten_nested(self, tmp_path):
        from undatum.tools import frequency
        from undatum.tools import schemas as tool_schemas

        src = _write_nested_jsonl(tmp_path / "nested.jsonl")
        result = frequency(str(src), "capital_city.lat", flatten_nested=True)
        assert result["ok"] is True
        assert result["data"]["total_rows"] == 2
        values = {item["value"] for item in result["data"]["top_values"]}
        assert "38.56" in values
        assert "42.87" in values
        for name in ("frequency", "deduplicate", "mask_fields", "sample_data"):
            schema = next(d for d in tool_schemas.UNDATUM_TOOL_DEFINITIONS if d["name"] == name)
            assert "flatten_nested" in schema["parameters"]["properties"]

    def test_tool_sample_data_flatten_nested(self, tmp_path):
        from undatum.tools import sample_data

        src = _write_nested_jsonl(tmp_path / "nested.jsonl")
        out = tmp_path / "out.jsonl"
        result = sample_data(str(src), str(out), n=2, confirm=True, flatten_nested=True)
        assert result["ok"] is True
        row = json.loads(out.read_text(encoding="utf8").splitlines()[0])
        assert "capital_city.lat" in row
        assert row["capital_city.lat"] in (38.56, 42.87)


class TestSdkAndCli:
    def test_dataset_read_preserves_table_on_write(self, tmp_path, monkeypatch):
        xlsx = _write_two_sheet_xlsx(tmp_path / "wb.xlsx")
        dst = tmp_path / "out.jsonl"
        seen = {}

        def fake_convert(self, fromfile, tofile, options, limit=None):
            seen["options"] = options
            Path(tofile).write_text('{"city":"Dushanbe"}\n', encoding="utf8")
            return None

        monkeypatch.setattr(Converter, "convert", fake_convert)
        Dataset.read(str(xlsx), table="Sheet2").write(str(dst))
        assert seen["options"]["table"] == "Sheet2"

    def test_convert_cli_help_exposes_wave1_flags(self):
        result = runner.invoke(app, ["convert", "--help"])
        assert result.exit_code == 0
        assert "--table" in result.stdout
        assert "--sheet" in result.stdout
        assert "--native-batch" in result.stdout
        assert "--no-native-batch" in result.stdout
        assert "--profile" in result.stdout
        assert "--columns" in result.stdout
        assert "--row-range" in result.stdout
        assert "--write-mode" in result.stdout
        assert "--row-group-size" in result.stdout
        assert "--trust" in result.stdout
        assert "--on-error" in result.stdout
        assert "--error-log" in result.stdout
        assert "--use-totals" in result.stdout
        assert "--filename-pattern" in result.stdout
        assert "--level" in result.stdout
        assert "--quotechar" in result.stdout

    def test_native_batch_default_is_none(self):
        default = inspect.signature(convert_cmd).parameters["native_batch"].default
        assert default is None

    def test_stats_and_schema_help_expose_flatten_nested(self):
        stats = runner.invoke(app, ["stats", "--help"])
        schema = runner.invoke(app, ["schema", "--help"])
        assert stats.exit_code == 0 and "--flatten-nested" in stats.stdout
        assert schema.exit_code == 0 and "--flatten-nested" in schema.stdout
        assert "--quotechar" in stats.stdout and "--quotechar" in schema.stdout
        assert "--table" in stats.stdout and "--table" in schema.stdout
        assert "--validate" in schema.stdout
        assert "--sample-size" in schema.stdout
        assert "--max-nested-depth" in stats.stdout and "--max-nested-depth" in schema.stdout
        assert "keep-nested" in stats.stdout and "keep-nested" in schema.stdout
        for cmd in (
            "analyze",
            "select",
            "head",
            "tail",
            "table",
            "uniq",
            "frequency",
            "headers",
            "sort",
            "sample",
            "search",
            "dedup",
            "fill",
            "rename",
            "mask",
            "plot",
            "validate",
            "sniff",
            "split",
            "join",
            "diff",
            "exclude",
        ):
            help_result = runner.invoke(app, [cmd, "--help"])
            assert help_result.exit_code == 0, help_result.stdout
            assert "--flatten-nested" in help_result.stdout
            assert "--max-nested-depth" in help_result.stdout
            assert "keep-nested" in help_result.stdout
        analyze_help = runner.invoke(app, ["analyze", "--help"])
        tail_help = runner.invoke(app, ["tail", "--help"])
        count_help = runner.invoke(app, ["count", "--help"])
        assert analyze_help.exit_code == 0 and "--table" in analyze_help.stdout
        assert tail_help.exit_code == 0 and "--table" in tail_help.stdout
        assert count_help.exit_code == 0 and "--table" in count_help.stdout
        assert "--quotechar" in analyze_help.stdout
        assert "--quotechar" in tail_help.stdout
        assert "--quotechar" in count_help.stdout
        for cmd in (
            "uniq",
            "frequency",
            "headers",
            "table",
            "doc",
            "sort",
            "dedup",
            "mask",
            "sample",
            "search",
            "sniff",
            "fill",
            "replace",
            "cat",
            "flatten",
            "slice",
            "enum",
            "reverse",
            "fixlengths",
            "rename",
            "explode",
            "transpose",
            "apply",
            "fmt",
            "join",
            "diff",
            "exclude",
            "plot",
            "validate",
            "split",
        ):
            help_result = runner.invoke(app, [cmd, "--help"])
            assert help_result.exit_code == 0, help_result.stdout
            assert "--table" in help_result.stdout
            assert "--quotechar" in help_result.stdout
        ingest_help = runner.invoke(app, ["ingest", "--help"])
        assert ingest_help.exit_code == 0
        assert "--source-table" in ingest_help.stdout
        assert "--quotechar" in ingest_help.stdout
        assert "--flatten-nested" in ingest_help.stdout
        assert "--max-nested-depth" in ingest_help.stdout
        assert "keep-nested" in ingest_help.stdout
        join_help = runner.invoke(app, ["join", "--help"])
        assert "--table2" in join_help.stdout
        for args in (
            ["package", "create", "--help"],
            ["package", "add-resource", "--help"],
        ):
            help_result = runner.invoke(app, args)
            assert help_result.exit_code == 0, help_result.stdout
            assert "--table" in help_result.stdout
            assert "--quotechar" in help_result.stdout
            assert "--flatten-nested" in help_result.stdout
            assert "--max-nested-depth" in help_result.stdout
            assert "keep-nested" in help_result.stdout
        schema_bulk_help = runner.invoke(app, ["schema-bulk", "--help"])
        assert schema_bulk_help.exit_code == 0
        assert "--table" in schema_bulk_help.stdout
        assert "--quotechar" in schema_bulk_help.stdout
        for args in (
            ["ai", "filter", "--help"],
            ["ai", "suggest", "--help"],
        ):
            help_result = runner.invoke(app, args)
            assert help_result.exit_code == 0, help_result.stdout
            assert "--table" in help_result.stdout
            assert "--quotechar" in help_result.stdout
        db_load_help = runner.invoke(app, ["db", "load", "--help"])
        assert db_load_help.exit_code == 0
        assert "--source-table" in db_load_help.stdout
        assert "--quotechar" in db_load_help.stdout
        assert "--flatten-nested" in db_load_help.stdout
        assert "--max-nested-depth" in db_load_help.stdout
        assert "keep-nested" in db_load_help.stdout
        doc_help = runner.invoke(app, ["doc", "--help"])
        assert doc_help.exit_code == 0
        assert "--flatten-nested" in doc_help.stdout
        assert "--max-nested-depth" in doc_help.stdout
        assert "keep-nested" in doc_help.stdout
        for args in (["tui", "--help"], ["web", "--help"]):
            help_result = runner.invoke(app, args)
            assert help_result.exit_code == 0, help_result.stdout
            assert "--table" in help_result.stdout
            assert "--quotechar" in help_result.stdout
            assert "--flatten-nested" in help_result.stdout
            assert "--max-nested-depth" in help_result.stdout
            assert "keep-nested" in help_result.stdout

    def test_schema_bulk_does_not_expose_flatten_nested(self):
        result = runner.invoke(app, ["schema-bulk", "--help"])
        assert result.exit_code == 0
        assert "--flatten-nested" not in result.stdout
